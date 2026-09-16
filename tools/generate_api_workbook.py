"""
Build `AI-GTM-OS-API-Acquisition.xlsx` — the working tracker for acquiring every
third-party API the platform can use.

    python tools/generate_api_workbook.py

Sheets:
    Start Here        how to use it, the phase legend, the 30-day plan
    API Master List   all providers: cost, difficulty, steps, .env keys + status tracker
    Acquisition Plan  phase roll-up (live formulas off the master list) + week-by-week
    Prerequisites     the non-API blockers that gate everything else
    Cost Model        indicative monthly spend by phase, with live formulas
    .env Key Map      every key in .env.example -> which provider fills it
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from api_catalog import COLUMNS, PREREQS, PROVIDERS  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "AI-GTM-OS-API-Acquisition.xlsx"

# ------------------------------------------------------------------ palette
INK = "FF0B0D13"
HEAD_BG = "FF1C2133"
HEAD_FG = "FFFFFFFF"
ACCENT = "FF7C6CFF"
BAND = "FFF6F5FF"
TITLE_FG = "FF3B2FA8"

PHASE_FILL = {
    "P0": PatternFill("solid", fgColor="FFFFE0E0"),
    "P1": PatternFill("solid", fgColor="FFFFEDD5"),
    "P2": PatternFill("solid", fgColor="FFE0F2FE"),
    "P3": PatternFill("solid", fgColor="FFE9E4FF"),
    "P4": PatternFill("solid", fgColor="FFF1F1F1"),
}
DIFF_FILL = {
    "Instant": "FFD9F2E3", "Easy": "FFE6F4D9", "Medium": "FFFFF2CC",
    "Hard": "FFFDE0DC", "Enterprise": "FFEFD9F7",
}
NEED_FILL = {
    "Required": "FFFFD9D9", "Recommended": "FFFFF0CC",
    "Optional": "FFEFEFEF", "Alternative": "FFE3ECFF",
}
STATUSES = ["Not started", "Prereq pending", "Applied", "In review", "Approved", "Live", "Blocked", "Skipped"]

thin = Side(style="thin", color="FFD7D7E5")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row=1, widths=None, height=34):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(bold=True, color=HEAD_FG, size=10)
        cell.fill = PatternFill("solid", fgColor=HEAD_BG)
        cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
        cell.border = BORDER
    ws.row_dimensions[row].height = height
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title_block(ws, title, subtitle, span=8):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=18, color=TITLE_FG)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    ws.row_dimensions[1].height = 28
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, color="FF555A70")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    ws.row_dimensions[2].height = 30


# =================================================================== Start Here
def sheet_start_here(wb):
    ws = wb.create_sheet("Start Here")
    ws.sheet_view.showGridLines = False
    title_block(ws, "AI GTM OS — API acquisition plan",
                "Every third-party API the platform can connect to, in the order you should actually get them. "
                "Work top-down: finish the Prerequisites, then Phase P0 → P1 → P2 → P3. "
                "Track progress in the Status column of 'API Master List' — the Acquisition Plan sheet rolls it up automatically.", span=6)

    rows = [
        ("", ""),
        ("THE ONE RULE", "Nothing here blocks you from using the product. Every provider left blank runs in SIMULATION mode, "
                         "so the whole system is demoable today. You are buying REAL DATA and REAL EXECUTION, one channel at a time."),
        ("", ""),
        ("START THESE ON DAY 1 (long lead times)", ""),
        ("1. Email sending domains + warmup", "2–3 WEEKS of warmup before you can send a single cold email. Longest pole in the project. "
                                              "Buy secondary domains today, create mailboxes, turn on warmup, then forget about it for 3 weeks."),
        ("2. Google Ads developer token", "Days to ~3 weeks of review. Apply the moment you have a Manager (MCC) account."),
        ("3. Meta Business Verification", "1–4 weeks. Unlocks Meta Ads advanced access, the Ad Library AND WhatsApp in one go."),
        ("4. LinkedIn Marketing API", "2–8 weeks and genuinely uncertain. Apply early; assume it may not land."),
        ("", ""),
        ("PHASE LEGEND", ""),
        ("P0 — Day 1", "A Groq API key. Free tier works for testing; every one of the 38 agents starts reasoning for real."),
        ("P1 — Week 1", "The wedge: Apollo + verifier + Instantly (outbound) · Google Ads (paid) · HubSpot (CRM) · Stripe + GA4 (revenue truth). "
                        "This is a complete, working revenue loop."),
        ("P2 — Week 2–4", "Second channel + research: Meta Ads, Firecrawl, Tavily, SerpApi, Search Console, Resend, Gmail, Cal.com."),
        ("P3 — Month 2–3", "Scale: LinkedIn, creative/video generation, YouTube, TikTok, WhatsApp, PostHog."),
        ("P4 — Later / optional", "Enterprise contracts (ZoomInfo, Semrush, Ahrefs) and alternatives you only need one of."),
        ("", ""),
        ("'ALTERNATIVE' MEANS PICK ONE", "Rows marked Alternative are pick-one-of-a-group: one CRM, one ESP, one sending platform, "
                                         "one email verifier, one SEO suite, one analytics tool. Do not buy the whole column."),
        ("", ""),
        ("MINIMUM VIABLE SPEND", "≈ $375/month of software gets the entire P0 + P1 loop running live on real data — "
                                 "Groq, Apollo, a verifier, Instantly, HubSpot, Stripe, GA4 and Google Ads. "
                                 "The ad budget itself is separate and will be the larger number. See the Cost Model sheet."),
        ("", ""),
        ("WHERE THE KEYS GO", "Either paste them into the repo-root .env file (see '.env Key Map'), or enter them in the app at "
                              "Integrations → click a provider → Save & test. Generate CREDENTIALS_ENCRYPTION_KEY first so "
                              "UI-entered keys are encrypted at rest."),
        ("", ""),
        ("SAFETY BEFORE YOU CONNECT ADS", "Settings → Control & guardrails: set the daily spend cap and leave "
                                          "'AI can launch paid campaigns without approval' OFF until you trust the decision feed."),
    ]
    r = 4
    for label, text in rows:
        ws.cell(r, 1, label).font = Font(bold=bool(label), size=11,
                                         color=TITLE_FG if label and label.isupper() else INK)
        c = ws.cell(r, 2, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(size=10)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        if text and len(text) > 110:
            ws.row_dimensions[r].height = 30
        r += 1

    ws.column_dimensions["A"].width = 38
    for col in "BCDEF":
        ws.column_dimensions[col].width = 22
    return ws


# =================================================================== Master list
def sheet_master(wb):
    ws = wb.create_sheet("API Master List")
    headers = [c[0] for c in COLUMNS]
    widths = [c[1] for c in COLUMNS]
    ws.append(headers)
    style_header(ws, 1, widths)

    for i, p in enumerate(PROVIDERS, start=1):
        ws.append([i, *p, "Not started", "", "", ""])

    last = ws.max_row
    ncols = len(headers)

    for row in ws.iter_rows(min_row=2, max_row=last, max_col=ncols):
        phase = row[1].value
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(size=9)
        row[0].alignment = Alignment(horizontal="center", vertical="top")
        row[1].fill = PHASE_FILL.get(phase, PatternFill())
        row[1].font = Font(size=9, bold=True)
        row[1].alignment = Alignment(horizontal="center", vertical="center")
        row[3].font = Font(size=10, bold=True)                      # Provider
        row[5].fill = PatternFill("solid", fgColor=NEED_FILL.get(row[5].value, "FFFFFFFF"))
        row[6].fill = PatternFill("solid", fgColor=DIFF_FILL.get(row[6].value, "FFFFFFFF"))
        row[10].number_format = '"$"#,##0'                          # Est. $/mo
        row[10].alignment = Alignment(horizontal="right", vertical="top")
        for idx in (12, 13):                                        # URLs
            cell = row[idx]
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(size=9, color="FF1155CC", underline="single")
        row[14].alignment = Alignment(vertical="top", wrap_text=True)   # steps
        row[19].number_format = "yyyy-mm-dd"                         # target date
        ws.row_dimensions[row[0].row].height = 96

    status_col = get_column_letter(headers.index("Status") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True, showDropDown=False)
    dv.error = "Pick a status from the list"
    ws.add_data_validation(dv)
    dv.add(f"{status_col}2:{status_col}{last}")

    rng = f"{status_col}2:{status_col}{last}"
    for value, colour, font_colour in [
        ("Live", "FFD1F0DB", "FF0B6B2E"), ("Approved", "FFE2F5CC", "FF3F6212"),
        ("In review", "FFFFF3C4", "FF8A5B00"), ("Applied", "FFE0EEFF", "FF1E40AF"),
        ("Blocked", "FFFFD6D6", "FFA11616"), ("Prereq pending", "FFF3E8FF", "FF6B21A8"),
        ("Skipped", "FFEDEDED", "FF777777"),
    ]:
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=[f'"{value}"'],
            fill=PatternFill("solid", fgColor=colour), font=Font(bold=True, size=9, color=font_colour)))

    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{last}"
    ws.freeze_panes = "E2"
    return ws, last


# =================================================================== Acquisition plan
def sheet_plan(wb, master_rows):
    ws = wb.create_sheet("Acquisition Plan")
    ws.sheet_view.showGridLines = False
    title_block(ws, "30-day acquisition plan",
                "Roll-up counts are live formulas over the Status column of 'API Master List' — update status there, "
                "these numbers follow.", span=8)

    m = "'API Master List'"
    ws.append([])
    ws.append(["Phase", "What it unlocks", "Items", "Live", "In flight", "Not started", "% complete", "Est. $/mo"])
    style_header(ws, ws.max_row, [10, 62, 9, 8, 11, 13, 13, 12], height=30)
    hdr = ws.max_row

    phase_desc = {
        "P0": "Agents reason for real instead of simulating",
        "P1": "Complete revenue loop: outbound + Google Ads + CRM + revenue truth",
        "P2": "Second paid channel, real research data, lifecycle email",
        "P3": "Scale: LinkedIn, video creative, more social, WhatsApp",
        "P4": "Enterprise contracts and pick-one alternatives",
    }
    r = hdr + 1
    for phase in ["P0", "P1", "P2", "P3", "P4"]:
        ws.cell(r, 1, phase).fill = PHASE_FILL[phase]
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 1).alignment = Alignment(horizontal="center")
        ws.cell(r, 2, phase_desc[phase]).alignment = Alignment(wrap_text=True, vertical="center")
        rng = f"{m}!$B$2:$B${master_rows}"
        st = f"{m}!$R$2:$R${master_rows}"
        cost = f"{m}!$K$2:$K${master_rows}"
        ws.cell(r, 3, f'=COUNTIF({rng},"{phase}")')
        ws.cell(r, 4, f'=COUNTIFS({rng},"{phase}",{st},"Live")')
        ws.cell(r, 5, f'=COUNTIFS({rng},"{phase}",{st},"Applied")+COUNTIFS({rng},"{phase}",{st},"In review")+COUNTIFS({rng},"{phase}",{st},"Approved")')
        ws.cell(r, 6, f'=COUNTIFS({rng},"{phase}",{st},"Not started")')
        ws.cell(r, 7, f'=IF(C{r}=0,0,D{r}/C{r})').number_format = "0%"
        ws.cell(r, 8, f'=SUMIFS({cost},{rng},"{phase}",{m}!$F$2:$F${master_rows},"Required")+'
                      f'SUMIFS({cost},{rng},"{phase}",{m}!$F$2:$F${master_rows},"Recommended")').number_format = '"$"#,##0'
        for c in range(1, 9):
            ws.cell(r, c).border = BORDER
            if c >= 3:
                ws.cell(r, c).alignment = Alignment(horizontal="center")
        ws.row_dimensions[r].height = 26
        r += 1

    ws.cell(r, 2, "TOTAL (Required + Recommended only)").font = Font(bold=True)
    for c, f in [(3, f"=SUM(C{hdr+1}:C{r-1})"), (4, f"=SUM(D{hdr+1}:D{r-1})"),
                 (5, f"=SUM(E{hdr+1}:E{r-1})"), (6, f"=SUM(F{hdr+1}:F{r-1})"),
                 (8, f"=SUM(H{hdr+1}:H{r-1})")]:
        cell = ws.cell(r, c, f)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER
    ws.cell(r, 7, f"=IF(C{r}=0,0,D{r}/C{r})").number_format = "0%"
    ws.cell(r, 7).font = Font(bold=True)
    ws.cell(r, 7).alignment = Alignment(horizontal="center")
    ws.cell(r, 8).number_format = '"$"#,##0'

    # ---- week by week
    r += 3
    ws.cell(r, 1, "WEEK BY WEEK").font = Font(bold=True, size=13, color=TITLE_FG)
    r += 1
    ws.cell(r, 1, "When")
    ws.cell(r, 2, "Do this")
    ws.cell(r, 3, "Why now")
    style_header(ws, r, height=24)
    for c in (1, 2, 3):
        ws.cell(r, c).border = BORDER
    ws.column_dimensions["C"].width = 52

    weeks = [
        ("Day 1 (morning)", "Buy 2–3 secondary sending domains, create 2–3 mailboxes each, set SPF/DKIM/DMARC, "
                            "sign up for Instantly, connect all mailboxes, switch warmup ON.",
         "Warmup takes 2–3 weeks of calendar time. Everything else can be done while it runs; this cannot be rushed."),
        ("Day 1 (afternoon)", "Create the Google Ads Manager (MCC) account and apply for the developer token. "
                              "Create the Meta Business Account and start Business Verification.",
         "Both are review queues measured in days-to-weeks. Get in the queue before you need them."),
        ("Day 1 (evening)", "Groq API key (free tier is fine to start). Generate CREDENTIALS_ENCRYPTION_KEY. "
                            "Set spend caps in Settings → Control & guardrails.",
         "Five minutes and all 38 agents stop simulating and start reasoning on your real business."),
        ("Day 2–3", "Apollo (paid tier for API) + one email verifier + HubSpot private app + Stripe test keys.",
         "This is the data spine: find → enrich → verify → CRM → revenue. Cheap and instant."),
        ("Day 3–5", "Firecrawl + Tavily. Run the Business Intelligence and Competitor Intelligence agents.",
         "≈$50 turns the knowledge graph from demo data into YOUR business, product, pricing and competitors. "
         "Highest leverage purchase in the list."),
        ("Week 1 end", "GA4 + Search Console service account. Slack bot for approvals.",
         "Closes the measurement loop and gets approval requests out of the dashboard and into your day."),
        ("Week 2", "Google Ads token should land → wire it up and launch one high-intent Search campaign. "
                   "Add Resend for lifecycle email. Add Cal.com for meeting booking.",
         "First real paid channel plus a booking path for the AI SDR."),
        ("Week 3", "Warmup finishes → launch the first real outbound sequence, throttled to 30–40/mailbox/day. "
                   "Meta verification should land → connect Meta Ads.",
         "Now both halves of the wedge (outbound + paid) are live and attributable."),
        ("Week 4", "Review the decision feed and attribution. Cut whatever is not working. "
                   "Apply for LinkedIn Marketing API if B2B enterprise is your motion.",
         "First real reallocation decision from the Autonomous Media Buyer, made on your own data."),
        ("Month 2–3", "Add creative generation (fal.ai / ElevenLabs), YouTube or TikTok if relevant, "
                      "PostHog for CRO, WhatsApp if your geo uses it.",
         "Only scale channels once the wedge has a proven CAC you are happy to spend more against."),
    ]
    r += 1
    for when, do, why in weeks:
        ws.cell(r, 1, when).font = Font(bold=True, size=10)
        ws.cell(r, 1).alignment = Alignment(vertical="top", wrap_text=True)
        for col, val in ((2, do), (3, why)):
            c = ws.cell(r, col, val)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(size=9)
        for c in (1, 2, 3):
            ws.cell(r, c).border = BORDER
        ws.row_dimensions[r].height = 56
        r += 1

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 62
    return ws


# =================================================================== Prerequisites
def sheet_prereqs(wb):
    ws = wb.create_sheet("Prerequisites")
    ws.sheet_view.showGridLines = False
    title_block(ws, "Prerequisites — do these before the APIs",
                "These are not APIs. They are the accounts, documents and decisions that gate the API applications. "
                "Skipping them is the single most common reason an acquisition plan stalls in week 2.", span=6)
    ws.append([])
    ws.append(["#", "Prerequisite", "Why it blocks you", "Owner", "Lead time", "How to do it", "Status"])
    style_header(ws, ws.max_row, [5, 40, 46, 16, 20, 70, 14], height=30)

    r = ws.max_row + 1
    for i, (item, why, owner, lead, how) in enumerate(PREREQS, start=1):
        ws.append([i, item, why, owner, lead, how, "Not started"])
        for c in range(1, 8):
            cell = ws.cell(r, c)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9)
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(r, 2).font = Font(size=10, bold=True)
        ws.row_dimensions[r].height = 62
        r += 1

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"G5:G{r-1}")
    ws.conditional_formatting.add(f"G5:G{r-1}", CellIsRule(
        operator="equal", formula=['"Live"'], fill=PatternFill("solid", fgColor="FFD1F0DB"),
        font=Font(bold=True, size=9, color="FF0B6B2E")))
    ws.freeze_panes = "A5"
    return ws


# =================================================================== Cost model
def sheet_cost(wb, master_rows):
    ws = wb.create_sheet("Cost Model")
    ws.sheet_view.showGridLines = False
    title_block(ws, "Indicative monthly cost",
                "List prices for small-team tiers, USD, early 2026 — ALWAYS verify on the vendor's pricing page before "
                "budgeting. Excludes the advertising budget itself, which is the biggest line item once you go live.", span=6)
    m = "'API Master List'"
    ws.append([])
    ws.append(["Scenario", "What's included", "Software $/mo", "Notes"])
    style_header(ws, ws.max_row, [34, 56, 16, 54], height=28)
    r = ws.max_row + 1

    scenarios = [
        ("Demo / simulation", "Nothing. Every provider blank.", 0,
         "The full product is usable and demoable — agents produce deterministic simulated output."),
        ("P0 — agents thinking", "Groq only", None,
         "The free tier covers testing. Add billing before you schedule agents — the free rate limits will throttle 38 of them."),
        ("P0+P1 — the wedge live", "Groq, Apollo, verifier, Instantly, HubSpot, Stripe, GA4, Google Ads, Slack", None,
         "MINIMUM VIABLE REAL SYSTEM. Complete revenue loop on real data. Add your ad budget on top."),
        ("+ P2 — second channel", "Meta Ads, Firecrawl, Tavily, SerpApi, GSC, Resend, Gmail, Cal.com", None,
         "Two paid channels, real research, lifecycle email."),
        ("+ P3 — scaling", "LinkedIn, creative generation, YouTube, TikTok, WhatsApp, PostHog", None,
         "LinkedIn/YouTube/TikTok/PostHog APIs are free — the increase here is creative generation (usage-based, "
         "so treat it as a floor). Only scale once the wedge shows a CAC you want to spend more against."),
        ("Everything incl. optional", "All Required + Recommended + Optional rows", None,
         "You will almost certainly never need this — 'Alternative' rows are pick-one."),
    ]
    formulas = {
        1: f'=SUMIFS({m}!$K$2:$K${master_rows},{m}!$B$2:$B${master_rows},"P0")',
        2: f'=SUMIFS({m}!$K$2:$K${master_rows},{m}!$B$2:$B${master_rows},"P0")+'
           f'SUMIFS({m}!$K$2:$K${master_rows},{m}!$B$2:$B${master_rows},"P1",{m}!$F$2:$F${master_rows},"Required")+'
           f'SUMIFS({m}!$K$2:$K${master_rows},{m}!$B$2:$B${master_rows},"P1",{m}!$F$2:$F${master_rows},"Recommended")',
        3: None, 4: None, 5: None,
    }
    for i, (name, includes, fixed, note) in enumerate(scenarios):
        ws.cell(r, 1, name).font = Font(bold=True, size=10)
        ws.cell(r, 2, includes).alignment = Alignment(wrap_text=True, vertical="center")
        if fixed is not None:
            ws.cell(r, 3, fixed)
        elif formulas.get(i):
            ws.cell(r, 3, formulas[i])
        else:
            prev = r - 1
            phase = {3: "P2", 4: "P3", 5: None}[i]
            if phase:
                # P3's value comes from Optional creative tools, so include Optional for that phase
                needs = ("Required", "Recommended") if phase == "P2" else ("Required", "Recommended", "Optional")
                terms = "+".join(
                    f'SUMIFS({m}!$K$2:$K${master_rows},{m}!$B$2:$B${master_rows},"{phase}",'
                    f'{m}!$F$2:$F${master_rows},"{n}")' for n in needs)
                ws.cell(r, 3, f"=C{prev}+{terms}")
            else:
                ws.cell(r, 3, f'=SUMIFS({m}!$K$2:$K${master_rows},{m}!$F$2:$F${master_rows},"Required")+'
                              f'SUMIFS({m}!$K$2:$K${master_rows},{m}!$F$2:$F${master_rows},"Recommended")+'
                              f'SUMIFS({m}!$K$2:$K${master_rows},{m}!$F$2:$F${master_rows},"Optional")')
        ws.cell(r, 3).number_format = '"$"#,##0'
        ws.cell(r, 3).font = Font(bold=True, size=11)
        ws.cell(r, 3).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(r, 4, note).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(r, 4).font = Font(size=9)
        for c in range(1, 5):
            ws.cell(r, c).border = BORDER
        ws.row_dimensions[r].height = 40
        r += 1

    r += 2
    ws.cell(r, 1, "COST NOTES").font = Font(bold=True, size=12, color=TITLE_FG)
    r += 1
    for note in [
        "Ad spend is NOT in these numbers. Budget it separately — it will dwarf the software cost.",
        "LLM usage scales with how often agents run. Start agents manually or daily, not hourly. "
        "Bulk work (10k personalised emails, reply classification) already routes to the cheaper GROQ_WORKER_MODEL.",
        "Apollo/ZoomInfo cost is driven by enrichment CREDITS, not seats. Verify before sending to avoid burning credits on bad data.",
        "Semrush and Ahrefs gate their APIs behind top-tier plans (~$500/mo each). SerpApi + Search Console cover "
        "most of the same ground for under $100.",
        "Video generation is the most volatile line item. Gate renders behind the Creative Testing Agent so you only "
        "render variants of proven winners.",
        "Set a hard daily spend cap in Settings → Control & guardrails BEFORE connecting any ad platform.",
    ]:
        c = ws.cell(r, 1, "•  " + note)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(size=9)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 28
        r += 1
    return ws


# =================================================================== env key map
def sheet_env_map(wb):
    ws = wb.create_sheet(".env Key Map")
    ws.sheet_view.showGridLines = False
    title_block(ws, ".env key map",
                "Every variable in .env.example, the section it lives in, and which provider fills it. "
                "Keys with no provider are platform settings you set yourself — no signup needed.", span=5)

    env_path = ROOT / ".env.example"
    sections: list[tuple[str, str]] = []
    current = "Core"
    if env_path.exists():
        for line in env_path.read_text(encoding="utf-8").splitlines():
            s = line.strip()
            sec = re.match(r"^#\s*(\d+\..+)$", s)
            if sec:
                current = sec.group(1).strip()
                continue
            kv = re.match(r"^([A-Z][A-Z0-9_]*)=", s)
            if kv:
                sections.append((current, kv.group(1)))

    key_to_provider: dict[str, str] = {}
    for p in PROVIDERS:
        provider, env_keys = p[2], p[14]   # (..., steps=13, env_keys=14, gotchas=15)
        for k in [x.strip() for x in env_keys.split(",")]:
            key_to_provider.setdefault(k, provider)

    # Keys in .env.example that have no dedicated row — minor alternatives, or values you invent yourself.
    key_to_provider.update({
        "META_WEBHOOK_VERIFY_TOKEN": "You invent this string (must match Meta)",
        "WHATSAPP_WEBHOOK_VERIFY_TOKEN": "You invent this string (must match Meta)",
        "META_AD_LIBRARY_ACCESS_TOKEN": "Meta Ads (usually the same token)",
        "GOOGLE_GEMINI_API_KEY": "Google AI Studio (optional 2nd model)",
        "BUFFER_ACCESS_TOKEN": "Buffer (alternative to Ayrshare)",
        "LUSHA_API_KEY": "Lusha (alternative enrichment)",
        "BRAVE_SEARCH_API_KEY": "Brave Search (alternative to Tavily)",
        "CLOUDFLARE_API_TOKEN": "Cloudflare (DNS for landing pages)",
        "CLOUDFLARE_ZONE_ID": "Cloudflare (DNS for landing pages)",
        "R2_ACCOUNT_ID": "Cloudflare R2 (storage)",
        "GOOGLE_CALENDAR_CLIENT_ID": "Google Cloud (same project as GA4)",
        "GOOGLE_CALENDAR_CLIENT_SECRET": "Google Cloud (same project as GA4)",
        "HUBSPOT_CLIENT_ID": "HubSpot (only for OAuth apps, not private apps)",
        "HUBSPOT_CLIENT_SECRET": "HubSpot (only for OAuth apps, not private apps)",
        "SMTP_HOST": "Your mailbox provider (native SMTP option)",
        "SMTP_PORT": "Your mailbox provider (native SMTP option)",
        "SMTP_USER": "Your mailbox provider (native SMTP option)",
        "SMTP_PASSWORD": "Your mailbox provider (native SMTP option)",
        "SMTP_USE_TLS": "Your mailbox provider (native SMTP option)",
        "BIGQUERY_PROJECT_ID": "Google Cloud (optional warehouse)",
        "BIGQUERY_DATASET": "Google Cloud (optional warehouse)",
        "SNOWFLAKE_ACCOUNT": "Snowflake (optional warehouse)",
        "SNOWFLAKE_USER": "Snowflake (optional warehouse)",
        "SNOWFLAKE_PASSWORD": "Snowflake (optional warehouse)",
        "HOTJAR_SITE_ID": "Hotjar (optional heatmaps)",
        "MICROSOFT_CLARITY_PROJECT_ID": "Microsoft Clarity (free heatmaps)",
        "GOOGLE_SERVICE_ACCOUNT_JSON": "Google Cloud service account (inline JSON)",
        "LLM_PROVIDER": "You choose: groq | anthropic | openai",
        "GROQ_BASE_URL": "Groq (tuning, leave blank)",
        "GROQ_TEMPERATURE": "Groq (tuning)",
        "GROQ_MAX_RETRIES": "Groq (tuning - raise on free tier)",
        "GROQ_TIMEOUT": "Groq (tuning)",
        "GROQ_SEND_REASONING_EFFORT": "Groq (tuning - model dependent)",
        "OPENAI_BASE_URL": "OpenAI-compatible gateway (optional)",
    })

    ws.append([])
    ws.append(["Section", ".env variable", "Filled by", "Type", "Required to run?"])
    style_header(ws, ws.max_row, [34, 44, 34, 22, 20], height=26)
    r = ws.max_row + 1

    platform_required = {"DJANGO_SECRET_KEY", "CREDENTIALS_ENCRYPTION_KEY"}
    for section, key in sections:
        provider = key_to_provider.get(key, "")
        kind = "Provider credential" if provider else ("Platform setting" if not key.startswith("VITE_") else "Frontend setting")
        req = "Yes" if key in platform_required else ("No — simulates if blank" if provider else "Has a default")
        ws.append([section, key, provider or "—", kind, req])
        for c in range(1, 6):
            cell = ws.cell(r, c)
            cell.border = BORDER
            cell.font = Font(size=9)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(r, 2).font = Font(size=9, name="Consolas")
        if provider:
            ws.cell(r, 3).font = Font(size=9, bold=True)
        else:
            ws.cell(r, 3).font = Font(size=9, color="FF999999")
        if key in platform_required:
            ws.cell(r, 5).fill = PatternFill("solid", fgColor="FFFFD9D9")
            ws.cell(r, 5).font = Font(size=9, bold=True, color="FFA11616")
        r += 1

    ws.auto_filter.ref = f"A4:E{r-1}"
    ws.freeze_panes = "A5"
    return ws, len(sections)


# =================================================================== build
def main():
    wb = Workbook()
    wb.remove(wb.active)

    sheet_start_here(wb)
    _, master_rows = sheet_master(wb)
    sheet_plan(wb, master_rows)
    sheet_prereqs(wb)
    sheet_cost(wb, master_rows)
    _, nkeys = sheet_env_map(wb)

    wb.properties.title = "AI GTM OS — API Acquisition Plan"
    wb.properties.creator = "AI GTM OS"
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"  {master_rows - 1} providers · {len(PREREQS)} prerequisites · {nkeys} .env keys mapped")


if __name__ == "__main__":
    main()
